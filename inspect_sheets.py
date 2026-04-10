#!/usr/bin/env python3
"""
Inspect Excel sheets to understand data structure
"""

from openpyxl import load_workbook

file_path = r"d:\Project\Excel\Small label\Small Label Printing.xlsx"
wb = load_workbook(file_path)

# Focus on key sheets
target_sheets = ['ProductLabel', 'PHQ', 'ExpiryDate', 'Abx', 'Extemporaneous Preparation']

for sheet_name in target_sheets:
    if sheet_name not in wb.sheetnames:
        print(f"Sheet '{sheet_name}' not found")
        continue

    print(f"\n{'='*60}")
    print(f"Sheet: {sheet_name}")
    print(f"{'='*60}")

    ws = wb[sheet_name]
    print(f"Dimensions: {ws.dimensions}")
    print(f"Max Row: {ws.max_row}, Max Col: {ws.max_column}")

    # Print first 5 rows
    print(f"\nFirst 5 rows:")
    for row_idx, row in enumerate(ws.iter_rows(max_row=5, values_only=True)):
        print(f"  Row {row_idx}: {row[:10]}")  # First 10 columns

    # Print all header columns
    print(f"\nAll headers (first row):")
    headers = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    for idx, header in enumerate(headers[:15]):
        if header:
            print(f"  Col {idx}: {header}")

    # Count non-empty rows (data rows)
    data_count = 0
    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row_idx > 0 and any(cell for cell in row):
            data_count += 1
    print(f"\nData rows (excluding header): {data_count}")
