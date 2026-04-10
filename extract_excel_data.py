#!/usr/bin/env python3
"""
Extract data from Small Label Printing.xlsx and create JSON backup
"""

import json
import sys
from datetime import datetime

def extract_from_openpyxl(file_path):
    """Extract data using openpyxl"""
    from openpyxl import load_workbook

    wb = load_workbook(file_path)
    sheets = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        data = []
        headers = []

        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if row_idx == 0:
                headers = [str(h).strip() if h else f'Column_{i}' for i, h in enumerate(row)]
            else:
                if any(cell is not None for cell in row):
                    row_dict = {}
                    for i, header in enumerate(headers):
                        row_dict[header] = row[i] if i < len(row) else None
                    data.append(row_dict)

        if data:
            sheets[sheet_name] = data

    return sheets

def normalize_drugs_from_phq(data):
    """Normalize drug/product data from PHQ sheet"""
    result = []
    for row in data:
        # PHQ sheet uses Column1, Column2, etc.
        col1 = str(row.get('Column1', '')).strip()
        col2 = str(row.get('Column2', '')).strip()
        col4 = str(row.get('Column4', '')).strip()  # Unit

        if not col1:
            continue

        result.append({
            'itemCode': col1.upper(),
            'name': col2,
            'unit': col4,
            'isSpecial': 'AUGM06' in col1,  # Special handling for augmented products
        })

    return result

def normalize_expiry_rules(data):
    """Normalize expiry rules data"""
    result = []
    for row in data:
        # ExpiryDate sheet
        item_code = str(row.get('ItemCode', '')).strip().upper()
        manufacturer = str(row.get('Manufacturer', '')).strip()

        if not item_code or not manufacturer:
            continue

        shelf_life_value = row.get('Expiry day', '')
        try:
            shelf_life_value = int(shelf_life_value) if shelf_life_value else ''
        except:
            shelf_life_value = str(shelf_life_value).strip() if shelf_life_value else ''

        result.append({
            'itemCode': item_code,
            'manufacturer': manufacturer,
            'shelfLifeValue': shelf_life_value,
            'shelfLifeUnit': str(row.get('Unit', 'Days')).strip() or 'Days',
            'storage': str(row.get('Storage Condition', '')).strip() or '',
            'warning': str(row.get('Warning #01', '')).strip() or '',
            'note': str(row.get('Warning #02', '')).strip() or '',
            'status': str(row.get('Status', 'Active')).strip() or 'Active',
        })

    return result

def normalize_abx(data):
    """Normalize ABX data"""
    result = []
    for row in data:
        item_code = str(row.get('Item Code', '')).strip().upper()
        if not item_code:
            continue

        result.append({
            'itemCode': item_code,
            'prefix': str(row.get('Prefix', '')).strip() or '',
            'packSize': '',
            'description': str(row.get('Item_description', '')).strip() or '',
        })

    return result

def main():
    file_path = r"d:\Project\Excel\Small label\Small Label Printing.xlsx"

    try:
        print("Extracting data from Excel file...")
        sheets = extract_from_openpyxl(file_path)
        print(f"Found {len(sheets)} sheets")
    except Exception as e:
        print(f"Error: {e}")
        sys.exit(1)

    # Create backup data structure
    backup_data = {
        'sml_drugs': [],
        'sml_expiryRules': [],
        'sml_abx': [],
        'sml_prepack': [],
        'sml_extemp': [],
        'sml_expAfterOpen': [],
        'sml_batchRecords': [],
        'sml_batchCounters': {'EC': 1, 'IPOP': 1},
        'sml_settings': {
            'location': 'IP/OP',
            'printerName': 'Thermal GX420t - PMS',
            'maintenancePin': 'NLTH',
            'showPrintDialog': True,
            'phqDatabasePath': '',
        },
    }

    # Process specific sheets
    if 'PHQ' in sheets:
        print("Processing PHQ sheet...")
        drugs = normalize_drugs_from_phq(sheets['PHQ'])
        backup_data['sml_drugs'].extend(drugs)
        print(f"  -> Extracted {len(drugs)} drugs/products")

    if 'ExpiryDate' in sheets:
        print("Processing ExpiryDate sheet...")
        rules = normalize_expiry_rules(sheets['ExpiryDate'])
        backup_data['sml_expiryRules'].extend(rules)
        print(f"  -> Extracted {len(rules)} expiry rules")

    if 'Abx' in sheets:
        print("Processing Abx sheet...")
        abx = normalize_abx(sheets['Abx'])
        backup_data['sml_abx'].extend(abx)
        print(f"  -> Extracted {len(abx)} ABX entries")

    # Add IDs to items
    timestamp = int(datetime.now().timestamp() * 1000)
    for category in ['sml_drugs', 'sml_expiryRules', 'sml_abx', 'sml_prepack', 'sml_extemp', 'sml_expAfterOpen']:
        for idx, item in enumerate(backup_data[category]):
            item['id'] = f"{timestamp}_{idx}"

    # Write JSON file
    output_file = r"d:\Project\Excel\Small label\small-label-backup.json"
    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump(backup_data, f, indent=2, ensure_ascii=False)

    print(f"\nBackup created: {output_file}")
    print(f"\nBackup Summary:")
    print(f"  - Drugs: {len(backup_data['sml_drugs'])}")
    print(f"  - Expiry Rules: {len(backup_data['sml_expiryRules'])}")
    print(f"  - ABX: {len(backup_data['sml_abx'])}")
    print("\nYou can import this JSON file in Settings → Data Backup & Restore")

if __name__ == '__main__':
    main()
